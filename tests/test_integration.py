"""
Prueba de integración de extremo a extremo (E2E).
"""

from pathlib import Path
import openpyxl
import pytest

from generador_606.config import AppConfig
from generar_606 import run_pipeline


def test_full_pipeline_run(tmp_path):
    out_dir = tmp_path / "salida_test"
    cfg = AppConfig.load(
        config_path="config.yaml",
        transacciones_path="muestra.xlsx",
        proveedores_path="Proveedores.xlsx",
        salida_dir=str(out_dir),
        tasa_retencion=0.10
    )

    exit_code = run_pipeline(cfg)
    assert exit_code == 0

    file_606 = out_dir / "formato_606.csv"
    file_excel = out_dir / "agrupacion_sin_comprobante.xlsx"
    file_incidencias = out_dir / "incidencias.txt"

    assert file_606.exists()
    assert file_excel.exists()
    assert file_incidencias.exists()

    # Verificar contenido de formato_606.csv
    with open(file_606, "r", encoding="utf-8-sig") as f:
        lines = [line.strip() for line in f if line.strip()]

    # 1 encabezado + 3 registros
    assert len(lines) == 4
    assert lines[1] == "06600056540,2,09,B0100000188,,202608,01,,,0.00,2033.90,2033.90,366.10,,0.00,0.00,366.10,,,,,,,,02"
    assert lines[2] == "06600056540,2,09,B0100000204,,202608,12,,,0.00,3389.83,3389.83,610.17,,0.00,0.00,610.17,,,,,,,,02"
    assert lines[3] == "132465555,1,09,B0100000383,,202608,17,,,0.00,26957.64,26957.64,3642.77,,0.00,0.00,3642.77,,,,,,,,02"

    # Verificar hojas de agrupacion_sin_comprobante.xlsx
    wb = openpyxl.load_workbook(file_excel, data_only=False)
    assert wb.sheetnames == ["Agrupado por proveedor", "Gastos menores B13", "Detalle"]

    ws_agrupado = wb["Agrupado por proveedor"]
    # Celda B3 es la tasa de retención
    assert ws_agrupado["B3"].value == 0.10

    # Total row formula en fila 9
    assert ws_agrupado["D9"].value == "=SUM(D6:D8)"
    assert ws_agrupado["J9"].value == "=SUM(J6:J8)"

    # Hoja Detalle tiene 6 filas de datos (filas 2 a 7)
    ws_detalle = wb["Detalle"]
    assert ws_detalle.max_row == 7

    wb.close()
